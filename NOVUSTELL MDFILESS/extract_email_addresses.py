#!/usr/bin/env python3
"""
Email Address Extractor for Novustell Travel
============================================

Extracts and analyzes email addresses from the cpanelemails directory
containing 146 email files spanning June 2021 to February 2024.

Features:
- Parses Maildir format email files
- Extracts From, To, Cc, Bcc, Reply-To addresses
- Categorizes emails by type
- Creates Excel report with detailed analysis
- Handles duplicate detection and frequency counting
"""

import os
import re
import email
import pandas as pd
from datetime import datetime
from collections import defaultdict, Counter
from email.header import decode_header
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

class EmailAddressExtractor:
    def __init__(self, cpanel_dir):
        self.cpanel_dir = cpanel_dir
        self.email_data = defaultdict(lambda: {
            'category': 'Unknown',
            'first_seen': None,
            'last_seen': None,
            'frequency': 0,
            'sources': set(),
            'domains': set(),
            'files': set()
        })
        
        # Email categorization patterns
        self.category_patterns = {
            'Business - Novustell': [
                r'.*@novustell\.com',
                r'.*@novustelltravel\.com',
                r'admin@novustell\.com',
                r'info@novustelltravel\.com'
            ],
            'Personal - Owner': [
                r'enockomondike@gmail\.com',
                r'enock.*@.*'
            ],
            'Google Services': [
                r'.*@google\.com',
                r'.*@gmail\.com',
                r'workspace-noreply@google\.com',
                r'.*\.bounces\.google\.com'
            ],
            'Travel Industry': [
                r'.*@katakenya\.org',
                r'communications@katakenya\.org',
                r'.*travel.*@.*',
                r'.*tourism.*@.*'
            ],
            'Marketing/Newsletter': [
                r'.*newsletter.*@.*',
                r'.*marketing.*@.*',
                r'.*@mailchimp\.com',
                r'.*@sendgrid\.net',
                r'bounces\+.*@.*'
            ],
            'System/Technical': [
                r'.*noreply.*@.*',
                r'.*no-reply.*@.*',
                r'.*@noc254\.com',
                r'.*@rcnoc\.com',
                r'postmaster@.*',
                r'mailer-daemon@.*'
            ]
        }
    
    def decode_email_header(self, header_value):
        """Decode email header that might be encoded"""
        if not header_value:
            return ""
        
        try:
            decoded_parts = decode_header(header_value)
            decoded_string = ""
            for part, encoding in decoded_parts:
                if isinstance(part, bytes):
                    decoded_string += part.decode(encoding or 'utf-8', errors='ignore')
                else:
                    decoded_string += str(part)
            return decoded_string
        except:
            return str(header_value)
    
    def extract_emails_from_header(self, header_value):
        """Extract email addresses from a header field"""
        if not header_value:
            return []
        
        # Decode the header first
        decoded_header = self.decode_email_header(header_value)
        
        # Email regex pattern - more comprehensive
        email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
        emails = re.findall(email_pattern, decoded_header, re.IGNORECASE)
        
        # Clean and validate emails
        cleaned_emails = []
        for email_addr in emails:
            email_addr = email_addr.lower().strip()
            # Basic validation
            if '@' in email_addr and '.' in email_addr.split('@')[1]:
                cleaned_emails.append(email_addr)
        
        return cleaned_emails
    
    def categorize_email(self, email_address):
        """Categorize email address based on patterns"""
        email_lower = email_address.lower()
        
        for category, patterns in self.category_patterns.items():
            for pattern in patterns:
                if re.match(pattern, email_lower, re.IGNORECASE):
                    return category
        
        # Additional logic for common domains
        domain = email_address.split('@')[1].lower()
        
        if domain in ['gmail.com', 'yahoo.com', 'hotmail.com', 'outlook.com']:
            return 'Personal - External'
        elif 'gov' in domain:
            return 'Government'
        elif 'edu' in domain:
            return 'Educational'
        elif any(word in domain for word in ['bank', 'finance', 'payment']):
            return 'Financial'
        else:
            return 'Business - External'
    
    def get_file_timestamp(self, filename):
        """Extract timestamp from Maildir filename"""
        try:
            timestamp = int(filename.split('.')[0])
            return datetime.fromtimestamp(timestamp)
        except:
            return datetime.now()
    
    def process_email_file(self, filepath):
        """Process a single email file and extract addresses"""
        try:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                msg = email.message_from_file(f)
            
            filename = os.path.basename(filepath)
            file_date = self.get_file_timestamp(filename)
            
            # Extract addresses from different header fields
            header_fields = {
                'From': msg.get('From', ''),
                'To': msg.get('To', ''),
                'Cc': msg.get('Cc', ''),
                'Bcc': msg.get('Bcc', ''),
                'Reply-To': msg.get('Reply-To', ''),
                'Return-Path': msg.get('Return-Path', ''),
                'Delivered-To': msg.get('Delivered-To', '')
            }
            
            for source, header_value in header_fields.items():
                if header_value:
                    emails = self.extract_emails_from_header(header_value)
                    
                    for email_addr in emails:
                        if email_addr:  # Skip empty emails
                            data = self.email_data[email_addr]
                            
                            # Update frequency
                            data['frequency'] += 1
                            
                            # Update dates
                            if data['first_seen'] is None or file_date < data['first_seen']:
                                data['first_seen'] = file_date
                            if data['last_seen'] is None or file_date > data['last_seen']:
                                data['last_seen'] = file_date
                            
                            # Add source and domain
                            data['sources'].add(source)
                            data['domains'].add(email_addr.split('@')[1])
                            data['files'].add(filename)
                            
                            # Categorize if not already done
                            if data['category'] == 'Unknown':
                                data['category'] = self.categorize_email(email_addr)
            
            return True
            
        except Exception as e:
            print(f"Error processing {filepath}: {e}")
            return False
    
    def process_all_files(self):
        """Process all email files in the cpanel directory"""
        print(f"🔍 Scanning {self.cpanel_dir} for email files...")
        
        email_files = []
        for filename in os.listdir(self.cpanel_dir):
            filepath = os.path.join(self.cpanel_dir, filename)
            if os.path.isfile(filepath) and not filename.startswith('.'):
                email_files.append(filepath)
        
        print(f"📧 Found {len(email_files)} email files to process")
        
        processed = 0
        failed = 0
        
        for i, filepath in enumerate(email_files, 1):
            print(f"Processing {i}/{len(email_files)}: {os.path.basename(filepath)}")
            
            if self.process_email_file(filepath):
                processed += 1
            else:
                failed += 1
        
        print(f"✅ Processing complete: {processed} successful, {failed} failed")
        return processed, failed
    
    def create_excel_report(self, output_file):
        """Create detailed Excel report of extracted email addresses"""
        print(f"📊 Creating Excel report: {output_file}")
        
        # Prepare data for DataFrame
        report_data = []
        
        for email_addr, data in self.email_data.items():
            report_data.append({
                'Email Address': email_addr,
                'Category': data['category'],
                'First Seen': data['first_seen'].strftime('%Y-%m-%d %H:%M:%S') if data['first_seen'] else '',
                'Last Seen': data['last_seen'].strftime('%Y-%m-%d %H:%M:%S') if data['last_seen'] else '',
                'Frequency': data['frequency'],
                'Sources': ', '.join(sorted(data['sources'])),
                'Domain': email_addr.split('@')[1],
                'Files Count': len(data['files'])
            })
        
        # Sort by frequency (descending) then by email address
        report_data.sort(key=lambda x: (-x['Frequency'], x['Email Address']))
        
        # Create DataFrame
        df = pd.DataFrame(report_data)
        
        # Create Excel file with formatting
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            # Main data sheet
            df.to_excel(writer, sheet_name='Email Addresses', index=False)
            
            # Get workbook and worksheet
            workbook = writer.book
            worksheet = writer.sheets['Email Addresses']
            
            # Format headers
            header_font = Font(bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
            
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='center')
            
            # Auto-adjust column widths
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width
            
            # Create summary sheet
            self.create_summary_sheet(writer, df)
        
        print(f"✅ Excel report created: {output_file}")
        return output_file
    
    def create_summary_sheet(self, writer, df):
        """Create summary analysis sheet"""
        # Category analysis
        category_counts = df['Category'].value_counts()
        domain_counts = df['Domain'].value_counts().head(20)
        
        # Create summary data
        summary_data = []
        summary_data.append(['NOVUSTELL TRAVEL - EMAIL ADDRESS ANALYSIS', ''])
        summary_data.append(['Report Generated', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
        summary_data.append(['', ''])
        summary_data.append(['OVERVIEW', ''])
        summary_data.append(['Total Unique Email Addresses', len(df)])
        summary_data.append(['Total Email Files Processed', len(set().union(*[data['files'] for data in self.email_data.values()]))]) 
        summary_data.append(['Date Range', f"{df['First Seen'].min()} to {df['Last Seen'].max()}"])
        summary_data.append(['', ''])
        summary_data.append(['CATEGORY BREAKDOWN', 'Count'])
        
        for category, count in category_counts.items():
            summary_data.append([category, count])
        
        summary_data.append(['', ''])
        summary_data.append(['TOP DOMAINS', 'Count'])
        
        for domain, count in domain_counts.items():
            summary_data.append([domain, count])
        
        # Write summary to sheet
        summary_df = pd.DataFrame(summary_data, columns=['Metric', 'Value'])
        summary_df.to_excel(writer, sheet_name='Summary', index=False)
        
        # Format summary sheet
        summary_ws = writer.sheets['Summary']
        
        # Format title
        summary_ws['A1'].font = Font(bold=True, size=14, color='366092')
        
        # Format section headers
        for row in summary_ws.iter_rows():
            if row[0].value in ['OVERVIEW', 'CATEGORY BREAKDOWN', 'TOP DOMAINS']:
                row[0].font = Font(bold=True, color='366092')
    
    def print_summary(self):
        """Print summary of extraction results"""
        total_emails = len(self.email_data)
        categories = defaultdict(int)
        domains = defaultdict(int)
        
        for email_addr, data in self.email_data.items():
            categories[data['category']] += 1
            domains[email_addr.split('@')[1]] += 1
        
        print(f"\n📊 EMAIL EXTRACTION SUMMARY")
        print(f"=" * 50)
        print(f"Total Unique Email Addresses: {total_emails}")
        print(f"Total Email Files Processed: {len(set().union(*[data['files'] for data in self.email_data.values()]))}")
        
        print(f"\n📂 CATEGORY BREAKDOWN:")
        for category, count in sorted(categories.items(), key=lambda x: x[1], reverse=True):
            percentage = (count / total_emails) * 100
            print(f"  {category}: {count} ({percentage:.1f}%)")
        
        print(f"\n🌐 TOP DOMAINS:")
        for domain, count in sorted(domains.items(), key=lambda x: x[1], reverse=True)[:10]:
            percentage = (count / total_emails) * 100
            print(f"  {domain}: {count} ({percentage:.1f}%)")


def main():
    """Main execution function"""
    print("🚀 Novustell Travel Email Address Extractor")
    print("=" * 50)
    
    # Configuration
    cpanel_dir = "/Users/djsean/Desktop/APPS2024/Novustellke/cpanelemails"
    output_file = "/Users/djsean/Desktop/APPS2024/Novustellke/novustell_email_addresses_extracted.xlsx"
    
    # Check if directory exists
    if not os.path.exists(cpanel_dir):
        print(f"❌ Error: Directory {cpanel_dir} not found")
        return
    
    # Initialize extractor
    extractor = EmailAddressExtractor(cpanel_dir)
    
    # Process all email files
    processed, failed = extractor.process_all_files()
    
    if processed == 0:
        print("❌ No email files were processed successfully")
        return
    
    # Create Excel report
    extractor.create_excel_report(output_file)
    
    # Print summary
    extractor.print_summary()
    
    print(f"\n✅ Email extraction complete!")
    print(f"📄 Excel report saved: {output_file}")
    print(f"📧 {len(extractor.email_data)} unique email addresses extracted")


if __name__ == "__main__":
    main()
