"""
Django management command to export Novustell Travel database to Excel format
"""

import os
from datetime import datetime
from django.core.management.base import BaseCommand, CommandError
from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from adminside.models import Destination, Package, Accommodation, Deal


class Command(BaseCommand):
    help = 'Export Novustell Travel database to Excel format for image management'

    def add_arguments(self, parser):
        parser.add_argument(
            '--output',
            type=str,
            default='novustell_travel_export.xlsx',
            help='Output Excel file name (default: novustell_travel_export.xlsx)'
        )
        parser.add_argument(
            '--path',
            type=str,
            default='.',
            help='Output directory path (default: current directory)'
        )

    def handle(self, *args, **options):
        output_file = options['output']
        output_path = options['path']
        
        # Create full file path
        full_path = os.path.join(output_path, output_file)
        
        self.stdout.write(
            self.style.SUCCESS(f'Starting export to: {full_path}')
        )

        try:
            # Create workbook
            wb = Workbook()
            
            # Remove default sheet
            wb.remove(wb.active)
            
            # Export each model
            self.export_destinations(wb)
            self.export_packages(wb)
            self.export_accommodations(wb)
            self.export_deals(wb)
            
            # Save workbook
            wb.save(full_path)
            
            self.stdout.write(
                self.style.SUCCESS(f'✅ Export completed successfully: {full_path}')
            )
            
            # Print summary
            self.print_summary(full_path)
            
        except Exception as e:
            raise CommandError(f'Export failed: {str(e)}')

    def export_destinations(self, workbook):
        """Export destinations data"""
        self.stdout.write('📍 Exporting destinations...')
        
        ws = workbook.create_sheet(title="Destinations")
        
        # Headers
        headers = [
            'ID', 'Name', 'Slug', 'Destination Type', 'Description', 
            'Current Image URL', 'Parent Destination', 'Meta Title', 
            'Meta Description', 'Starting Price', 'Is Active', 
            'Created At', 'Updated At'
        ]
        
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        
        # Write data
        destinations = Destination.objects.all().order_by('id')
        for row, dest in enumerate(destinations, 2):
            ws.cell(row=row, column=1, value=dest.id)
            ws.cell(row=row, column=2, value=dest.name)
            ws.cell(row=row, column=3, value=dest.slug)
            ws.cell(row=row, column=4, value=dest.get_destination_type_display())
            ws.cell(row=row, column=5, value=self.clean_text(dest.description))
            ws.cell(row=row, column=6, value=self.get_image_url(dest.image))
            ws.cell(row=row, column=7, value=dest.parent.name if dest.parent else '')
            ws.cell(row=row, column=8, value=dest.meta_title)
            ws.cell(row=row, column=9, value=dest.meta_description)
            ws.cell(row=row, column=10, value=float(dest.starting_price) if dest.starting_price else '')
            ws.cell(row=row, column=11, value='Yes' if dest.is_active else 'No')
            ws.cell(row=row, column=12, value=dest.created_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=13, value=dest.updated_at.strftime('%Y-%m-%d %H:%M:%S'))
        
        # Auto-adjust column widths
        self.adjust_column_widths(ws)
        
        self.stdout.write(f'   ✅ Exported {destinations.count()} destinations')

    def export_packages(self, workbook):
        """Export packages data"""
        self.stdout.write('📦 Exporting packages...')
        
        ws = workbook.create_sheet(title="Packages")
        
        # Headers
        headers = [
            'ID', 'Name', 'Slug', 'Description', 'Main Destination', 
            'Adult Price', 'Child Price', 'Duration Days', 'Status', 
            'Featured Image URL', 'Is Featured', 'Created At', 'Updated At'
        ]
        
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="FF9D00", end_color="FF9D00", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        
        # Write data
        packages = Package.objects.all().select_related('main_destination').order_by('id')
        for row, pkg in enumerate(packages, 2):
            ws.cell(row=row, column=1, value=pkg.id)
            ws.cell(row=row, column=2, value=pkg.name)
            ws.cell(row=row, column=3, value=pkg.slug)
            ws.cell(row=row, column=4, value=self.clean_text(pkg.description))
            ws.cell(row=row, column=5, value=pkg.main_destination.name if pkg.main_destination else '')
            ws.cell(row=row, column=6, value=float(pkg.adult_price) if pkg.adult_price else '')
            ws.cell(row=row, column=7, value=float(pkg.child_price) if pkg.child_price else '')
            ws.cell(row=row, column=8, value=pkg.duration_days)
            ws.cell(row=row, column=9, value=pkg.get_status_display())
            ws.cell(row=row, column=10, value=self.get_image_url(pkg.featured_image))
            ws.cell(row=row, column=11, value='Yes' if pkg.is_featured else 'No')
            ws.cell(row=row, column=12, value=pkg.created_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=13, value=pkg.updated_at.strftime('%Y-%m-%d %H:%M:%S'))
        
        # Auto-adjust column widths
        self.adjust_column_widths(ws)
        
        self.stdout.write(f'   ✅ Exported {packages.count()} packages')

    def export_accommodations(self, workbook):
        """Export accommodations data"""
        self.stdout.write('🏨 Exporting accommodations...')
        
        ws = workbook.create_sheet(title="Accommodations")
        
        # Headers
        headers = [
            'ID', 'Name', 'Slug', 'Description', 'Destination',
            'Accommodation Type', 'Price Per Room Per Night', 'Featured Image URL',
            'Is Active', 'Created At', 'Updated At'
        ]
        
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="28A745", end_color="28A745", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        
        # Write data
        accommodations = Accommodation.objects.all().select_related('destination').order_by('id')
        for row, acc in enumerate(accommodations, 2):
            ws.cell(row=row, column=1, value=acc.id)
            ws.cell(row=row, column=2, value=acc.name)
            ws.cell(row=row, column=3, value=acc.slug)
            ws.cell(row=row, column=4, value=self.clean_text(acc.description))
            ws.cell(row=row, column=5, value=acc.destination.name if acc.destination else '')
            ws.cell(row=row, column=6, value=acc.get_accommodation_type_display())
            ws.cell(row=row, column=7, value=float(acc.price_per_room_per_night) if acc.price_per_room_per_night else '')
            ws.cell(row=row, column=8, value=self.get_image_url(acc.image))
            ws.cell(row=row, column=9, value='Yes' if acc.is_active else 'No')
            ws.cell(row=row, column=10, value=acc.created_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=11, value=acc.updated_at.strftime('%Y-%m-%d %H:%M:%S'))
        
        # Auto-adjust column widths
        self.adjust_column_widths(ws)
        
        self.stdout.write(f'   ✅ Exported {accommodations.count()} accommodations')

    def export_deals(self, workbook):
        """Export deals data"""
        self.stdout.write('💰 Exporting deals...')
        
        ws = workbook.create_sheet(title="Deals")
        
        # Headers
        headers = [
            'ID', 'Title', 'Slug', 'Description', 'Discount Percentage', 
            'Original Price', 'Discounted Price', 'Valid From', 'Valid Until', 
            'Featured Image URL', 'Is Active', 'Is Featured', 'Created At', 'Updated At'
        ]
        
        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DC3545", end_color="DC3545", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        
        # Write data
        deals = Deal.objects.all().order_by('id')
        for row, deal in enumerate(deals, 2):
            ws.cell(row=row, column=1, value=deal.id)
            ws.cell(row=row, column=2, value=deal.title)
            ws.cell(row=row, column=3, value=deal.slug)
            ws.cell(row=row, column=4, value=self.clean_text(deal.description))
            ws.cell(row=row, column=5, value=float(deal.discount_percentage))
            ws.cell(row=row, column=6, value=float(deal.original_price))
            ws.cell(row=row, column=7, value=float(deal.discounted_price))
            ws.cell(row=row, column=8, value=deal.valid_from.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=9, value=deal.valid_until.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=10, value=self.get_image_url(deal.featured_image))
            ws.cell(row=row, column=11, value='Yes' if deal.is_active else 'No')
            ws.cell(row=row, column=12, value='Yes' if deal.is_featured else 'No')
            ws.cell(row=row, column=13, value=deal.created_at.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=14, value=deal.updated_at.strftime('%Y-%m-%d %H:%M:%S'))
        
        # Auto-adjust column widths
        self.adjust_column_widths(ws)
        
        self.stdout.write(f'   ✅ Exported {deals.count()} deals')

    def clean_text(self, text):
        """Clean HTML tags and excessive whitespace from text"""
        if not text:
            return ''
        
        # Remove HTML tags (basic cleaning)
        import re
        clean = re.compile('<.*?>')
        text = re.sub(clean, '', str(text))
        
        # Remove excessive whitespace
        text = ' '.join(text.split())
        
        return text[:500]  # Limit to 500 characters for Excel

    def get_image_url(self, image_field):
        """Extract image URL from UploadCare field"""
        if not image_field:
            return ''
        
        try:
            if hasattr(image_field, 'cdn_url'):
                return image_field.cdn_url
            elif hasattr(image_field, 'url'):
                return image_field.url
            else:
                return str(image_field)
        except:
            return ''

    def adjust_column_widths(self, worksheet):
        """Auto-adjust column widths based on content"""
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            # Set width with some padding, max 50 characters
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width

    def print_summary(self, file_path):
        """Print export summary"""
        file_size = os.path.getsize(file_path)
        file_size_mb = file_size / (1024 * 1024)
        
        self.stdout.write('\n' + '='*60)
        self.stdout.write(self.style.SUCCESS('📊 EXPORT SUMMARY'))
        self.stdout.write('='*60)
        self.stdout.write(f'📁 File: {file_path}')
        self.stdout.write(f'📏 Size: {file_size_mb:.2f} MB')
        self.stdout.write(f'📅 Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
        self.stdout.write('\n📋 Sheets created:')
        self.stdout.write('   • Destinations - All destination data with image URLs')
        self.stdout.write('   • Packages - All package data with image URLs')
        self.stdout.write('   • Accommodations - All accommodation data with image URLs')
        self.stdout.write('   • Deals - All deal data with image URLs')
        self.stdout.write('\n💡 Usage for bulk image updates:')
        self.stdout.write('   1. Open the Excel file')
        self.stdout.write('   2. Add new image URLs in the image URL columns')
        self.stdout.write('   3. Use the bulk_update_images command to apply changes')
        self.stdout.write('='*60)
